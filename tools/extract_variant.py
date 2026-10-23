"""Liest eine Excel-Variante (eigene Eingaben + Ergebnisse) als zusätzlichen Testfall ein.

Aufruf:  python tools/extract_variant.py NAME PFAD/ZUR/EXCEL.xlsx
Erzeugt: tests/reference/varianten/NAME_eingaben.json, NAME_zellwerte.json, NAME_stunden.csv.gz
"""
from __future__ import annotations

import gzip
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from lcoh.input_spec import SPECS  # noqa: E402

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parent.parent
SHEETS = {"Blatt 2": "2. Ein- und Ausgaben", "Blatt 5": "5. Strompreis", "Blatt 8": "8. Marktpreise"}


def main(name: str, path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    inputs = {}
    for s in SPECS:
        m = re.match(r"(Blatt \d), ([A-Z]+\d+)", s.excel)
        if not m or m.group(1) not in SHEETS:
            continue
        v = wb[SHEETS[m.group(1)]][m.group(2)].value
        if v is None:
            continue
        inputs[s.key] = v
    out = ROOT / "tests" / "reference" / "varianten"
    out.mkdir(parents=True, exist_ok=True)
    (out / f"{name}_eingaben.json").write_text(json.dumps(inputs, ensure_ascii=False, indent=1), encoding="utf8")

    ref = {}
    for sheet, short in (("2. Ein- und Ausgaben", "B2"), ("3. Nebenrechnungen", "B3"),
                         ("5. Strompreis", "B5"), ("6. Stacktausch", "B6")):
        for row in wb[sheet].iter_rows():
            for c in row:
                if isinstance(c.value, (int, float, bool)):
                    ref[f"{short}!{c.coordinate}"] = c.value
    ws4 = wb["4. Bezüge und Verbräuche"]
    for row in ws4.iter_rows(min_row=1, max_row=13):
        for c in row:
            if isinstance(c.value, (int, float, bool)):
                ref[f"B4!{c.coordinate}"] = c.value
    (out / f"{name}_zellwerte.json").write_text(json.dumps(ref, ensure_ascii=False), encoding="utf8")

    from openpyxl.utils import get_column_letter, column_index_from_string as ci
    cols = [get_column_letter(i) for i in range(ci("C"), ci("AJ") + 1)]
    rows = {c: [ws4[f"{c}{14 + i}"].value for i in range(8760)] for c in cols}
    pd.DataFrame(rows).to_csv(out / f"{name}_stunden.csv.gz", index=False, compression="gzip")
    print(name, len(inputs), "Eingaben,", len(ref), "Zellen")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
